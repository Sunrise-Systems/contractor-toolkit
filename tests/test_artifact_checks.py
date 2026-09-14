"""Synthetic artifact readback regression tests; no client files."""
import hashlib
import importlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch


def verify(path, expectations):
    try:
        module = importlib.import_module('scripts.artifact_checks')
    except ModuleNotFoundError:
        raise AssertionError('artifact verification API is missing')
    return module.verify_artifact(path, expectations)


class ArtifactTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.fields = {'project_id': 'SYN-1', 'company': 'Synthetic Builder', 'revision': 'R1'}
        self.exp = {'schema_version': 1, 'approved_fields': self.fields,
                    'required_text': ['Scope'], 'print_css': '@media print { body { color: black; } }'}
    def html(self, extra=''):
        p = self.root / 'bid.html'
        body = ''.join(f'<p data-field="{k}">{v}</p>' for k, v in self.fields.items())
        p.write_text(f'<html><head><style>{self.exp["print_css"]}</style></head>'
                     f'<body>{body}<h1>Scope</h1>{extra}</body></html>', encoding='utf-8')
        return p
    def test_html_rejects_submission_elements(self):
        for tag in ('form', 'input', 'button', 'select', 'textarea'):
            with self.subTest(tag=tag), self.assertRaisesRegex(ValueError, 'submission'):
                verify(self.html(f'<{tag}></{tag}>'), self.exp)
    def test_html_rejects_submission_attributes(self):
        for attr in ('action', 'formaction'):
            for value in ('https://invalid.test/submit', '/submit', '', '#scope', 'mailto:a@invalid.test'):
                with self.subTest(attr=attr, value=value), self.assertRaisesRegex(ValueError, 'submission'):
                    verify(self.html(f'<div {attr}="{value}">Review</div>'), self.exp)
    def test_closed_dialog_cannot_certify_a_visible_revision_mismatch(self):
        p = self.html()
        p.write_text(p.read_text().replace('<p data-field="revision">R1</p>',
                     '<dialog><p data-field="revision">R1</p></dialog><p>revision: R2</p>'))
        with self.assertRaisesRegex(ValueError, 'hidden'):
            verify(p, self.exp)
    def test_closed_popover_cannot_certify_a_visible_revision_mismatch(self):
        for attr in ('popover', 'popover="auto"', 'popover="manual"'):
            with self.subTest(attr=attr):
                p = self.html()
                p.write_text(p.read_text().replace('<p data-field="revision">R1</p>',
                             f'<div {attr}><p data-field="revision">R1</p></div>'
                             '<p>revision: R2</p>'))
                with self.assertRaisesRegex(ValueError, 'hidden'):
                    verify(p, self.exp)
    def test_html_invisible_subtrees_cannot_supply_required_text(self):
        for wrapper in ('<template>{}</template>', '<div inert>{}</div>',
                        '<div aria-hidden="TRUE">{}</div>', '<div hidden>{}</div>',
                        '<details><summary>Heading</summary>{}</details>',
                        '<details><summary>Heading</summary><summary>{}</summary></details>'):
            p = self.html()
            p.write_text(p.read_text().replace('<h1>Scope</h1>', wrapper.format('<h1>Scope</h1>')))
            with self.subTest(wrapper=wrapper), self.assertRaises(ValueError):
                verify(p, self.exp)
        for wrapper in ('<dialog open>{}</dialog>', '<details open>{}</details>',
                        '<details><summary>{}</summary></details>', '<div aria-hidden="false">{}</div>'):
            p = self.html()
            p.write_text(p.read_text().replace('<h1>Scope</h1>', wrapper.format('<h1>Scope</h1>')))
            self.assertEqual(verify(p, self.exp)['checks']['readback'], 'passed')
    def test_template_is_explicitly_rejected(self):
        with self.assertRaisesRegex(ValueError, 'hidden'):
            verify(self.html('<template><p>Unrendered content</p></template>'), self.exp)
    def test_hidden_css_is_rejected_even_in_canonical_print_css(self):
        for declaration in ('display: none', 'visibility: hidden', 'visibility: collapse',
                            'content-visibility: hidden', 'opacity: 0', 'display:/**/none'):
            for canonical in (True, False):
                with self.subTest(css=declaration, canonical=canonical):
                    self.exp['print_css'] = '@media print { body { color: black; } }'
                    if canonical:
                        self.exp['print_css'] = '@media print { body { ' + declaration + '; } }'
                    p = self.html('' if canonical else f'<div style="{declaration}">Hidden</div>')
                    with self.assertRaisesRegex(ValueError, 'hidden.*CSS'):
                        verify(p, self.exp)
    def test_duplicate_attributes_cannot_hide_active_content(self):
        for extra in ('<a href="javascript:alert(1)" href="#scope">Review</a>',
                      '<meta http-equiv="refresh" http-equiv="x" content="0;url=https://invalid.test/">'):
            with self.subTest(extra=extra), self.assertRaises(ValueError):
                verify(self.html(extra), self.exp)
    def test_html_receipt_binds_saved_bytes_and_semantics(self):
        p = self.html()
        receipt = verify(p, self.exp)
        self.assertEqual(receipt['file_sha256'], hashlib.sha256(p.read_bytes()).hexdigest())
        semantic = hashlib.sha256(json.dumps(self.fields, sort_keys=True, separators=(',', ':'),
                                            ensure_ascii=False).encode()).hexdigest()
        self.assertEqual(receipt['semantic_sha256'], semantic)
        self.assertEqual(receipt['path'], str(p.resolve()))
        self.assertEqual(receipt['issuance'], 'not_issued')
        self.assertIn('visual_layout_not_inspected', receipt['inspection_gaps'])
        self.assertEqual(receipt['project_id'], 'SYN-1')
        self.assertTrue(receipt['parser_version'])
        self.assertTrue(receipt['timestamp'])
    def test_html_rejects_unapproved_or_incomplete_content(self):
        for old, new in [('SYN-1', 'SYN-10'), ('Scope', 'Other'),
                         (self.exp['print_css'], ''), ('R1', '{{REVISION}}')]:
            with self.subTest(old=old):
                p = self.html()
                p.write_text(p.read_text().replace(old, new))
                with self.assertRaises(ValueError):
                    verify(p, self.exp)
        for extra in ['<script>alert(1)</script>', '<img src="https://invalid.test/logo">',
                      '<p data-field="project_id">OTHER</p>', '<p>{{LOGO}}</p>']:
            with self.subTest(extra=extra), self.assertRaises(ValueError):
                verify(self.html(extra), self.exp)
    def test_html_rejects_active_url_attributes(self):
        for extra in ['<a href="javascript:alert(1)">Review</a>',
                      '<a href=" &#x09;JaVa&#x0A;ScRiPt:alert(1)">Review</a>',
                      '<a href="vbscript:msgbox(1)">Review</a>',
                      '<a href="data:text/html,active">Review</a>',
                      '<form action="javascript:alert(1)"></form>',
                      '<button formaction="javascript:alert(1)">Review</button>',
                      '<svg><a xlink:href="javascript:alert(1)">Review</a></svg>']:
            with self.subTest(extra=extra), self.assertRaisesRegex(ValueError, 'active URL'):
                verify(self.html(extra), self.exp)
        for href in ['https://invalid.test/review', 'mailto:review@invalid.test', '#scope', 'review.html']:
            with self.subTest(href=href):
                self.assertEqual(verify(self.html(f'<a href="{href}">Review</a>'), self.exp)
                                 ['checks']['readback'], 'passed')
    def test_html_rejects_meta_refresh(self):
        for directive in ['refresh', 'ReFrEsH']:
            for content in ['0;url=https://invalid.test/', '0']:
                with self.subTest(directive=directive, content=content):
                    p = self.html()
                    p.write_text(p.read_text().replace('<head>',
                        f'<head><meta http-equiv="{directive}" content="{content}">'))
                    with self.assertRaisesRegex(ValueError, 'refresh'):
                        verify(p, self.exp)
    def test_stale_receipt_and_invalid_expectations_fail_closed(self):
        p = self.html()
        receipt = verify(p, self.exp)
        for changes in [{'file_sha256': '0' * 64}, {'semantic_sha256': '0' * 64},
                        {'schema_version': 2}, {'unexpected': True}, {'approved_fields': {}}]:
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                verify(p, dict(self.exp, **changes))
        p.write_text(p.read_text() + '\n')
        with self.assertRaises(ValueError):
            verify(p, dict(self.exp, file_sha256=receipt['file_sha256']))
        with self.assertRaises(ValueError):
            verify(self.root / 'missing.html', self.exp)
        with patch.dict('sys.modules', {'bs4': None}), self.assertRaises(ValueError):
            verify(p, self.exp)
    def test_docx_reads_tables_headers_footers_and_split_tokens(self):
        from docx import Document
        p = self.root / 'bid.docx'
        doc = Document()
        doc.add_paragraph('Scope')
        doc.add_paragraph('project_id: SYN-1')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'company: Synthetic Builder'
        doc.sections[0].header.paragraphs[0].text = 'revision: R1'
        doc.sections[0].footer.paragraphs[0].text = 'amount: 100.00'
        doc.save(p)
        exp = dict(self.exp, approved_fields=dict(self.fields, amount='100.00'))
        self.assertEqual(verify(p, exp)['parser_version'], '1.2.0')
        doc.sections[0].footer.paragraphs[0].text = 'amount: 1000.00'
        doc.save(p)
        with self.assertRaises(ValueError):
            verify(p, exp)
        doc.sections[0].footer.paragraphs[0].text = 'amount: 100.00'
        para = doc.sections[0].first_page_header.paragraphs[0]
        para.add_run('{{TO'); para.add_run('KEN}}')
        doc.save(p)
        with self.assertRaises(ValueError):
            verify(p, exp)
        p.write_bytes(b'corrupt zip')
        with self.assertRaises(ValueError):
            verify(p, exp)
    def workbook(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = 'Estimate'
        for i, value in enumerate(self.fields.values(), 1):
            sheet.cell(i, 1, value)
        sheet['A4'] = 'Scope'
        sheet['B1'], sheet['B2'], sheet['B3'] = 2, 10, 20
        sheet['A5'] = 'APP-0'
        p = self.root / 'bid.xlsx'
        book.save(p)
        exp = dict(self.exp, field_cells=dict(zip(self.fields, ['Estimate!A1', 'Estimate!A2', 'Estimate!A3'])),
                   required_sheets=['Estimate'], cells={'Estimate!B3': 20, 'Estimate!A5': 'APP-0'})
        return p, book, exp
    def test_xlsx_exact_cells_identity_tokens_and_prior_link(self):
        p, book, exp = self.workbook()
        self.assertEqual(verify(p, exp)['parser_version'], '3.1.5')
        for cell, value in [('B3', 21), ('A5', 'APP-WRONG'), ('A1', 'SYN-10'), ('Z9', '{{TOKEN}}')]:
            with self.subTest(cell=cell):
                p, book, exp = self.workbook()
                book.active[cell] = value
                book.save(p)
                with self.assertRaises(ValueError):
                    verify(p, exp)
        with self.assertRaises(ValueError):
            verify(p, dict(exp, required_sheets=['Missing']))
    def test_xlsx_cell_scalars_distinguish_booleans_from_numbers(self):
        p, book, exp = self.workbook()
        for actual, expected in [(True, 1), (False, 0), (1, True), (0, False),
                                 (True, 1.0), (False, 0.0)]:
            with self.subTest(actual=actual, expected=expected):
                book.active['B3'] = actual
                book.save(p)
                exp['cells']['Estimate!B3'] = expected
                with self.assertRaisesRegex(ValueError, 'cell mismatch: Estimate!B3'):
                    verify(p, exp)
        for actual, expected in [(True, True), (False, False), (1, 1.0),
                                 (0.0, 0), ('1', '1'), (None, None)]:
            with self.subTest(actual=actual, expected=expected):
                book.active['B3'] = actual
                book.save(p)
                exp['cells']['Estimate!B3'] = expected
                self.assertEqual(verify(p, exp)['checks']['readback'], 'passed')
    def test_xlsx_formula_text_or_cache_never_proves_financial_result(self):
        for formula in ['=B1*B2', '=SUM(B1:B2)', '=WEBSERVICE("https://invalid.test")']:
            with self.subTest(formula=formula):
                p, book, exp = self.workbook()
                book.active['Z1'] = formula
                book.save(p)
                with self.assertRaisesRegex(ValueError, 'formula'):
                    verify(p, exp)
    def test_financial_crossfoot_recomputed_with_explicit_decimal_rounding(self):
        p, book, exp = self.workbook()
        book.active['B4'], book.active['B5'], book.active['B6'] = 3, 23, 20
        book.save(p)
        exp['calculations'] = [
            {'target': 'Estimate!B3', 'operation': 'product', 'inputs': ['Estimate!B1', 'Estimate!B2'],
             'quantum': '0.01', 'rounding': 'ROUND_HALF_UP'},
            {'target': 'Estimate!B5', 'operation': 'sum', 'inputs': ['Estimate!B3', 'Estimate!B4'],
             'quantum': '0.01', 'rounding': 'ROUND_HALF_UP'},
            {'target': 'Estimate!B6', 'operation': 'subtract', 'inputs': ['Estimate!B5', 'Estimate!B4'],
             'quantum': '0.01', 'rounding': 'ROUND_HALF_UP'}]
        self.assertEqual(verify(p, exp)['checks']['calculations'], 'passed')
        for value in [9, 'NaN', True]:
            book.active['B2'] = value
            book.save(p)
            with self.subTest(value=value), self.assertRaises(ValueError):
                verify(p, exp)
    def test_pdf_cannot_reuse_source_success_without_text_and_layout_proof(self):
        p = self.html().rename(self.root / 'bid.pdf')
        with patch.dict('sys.modules', {'pypdf': None}), self.assertRaisesRegex(ValueError, 'PDF.*unavailable'):
            verify(p, self.exp)
        p.write_bytes(b'%PDF-1.7\ncorrupt')
        with self.assertRaisesRegex(ValueError, 'PDF'):
            verify(p, self.exp)
    def test_html_hidden_fields_and_active_resources_do_not_verify(self):
        for extra in ['<style>p { display: none }</style>', '<p onclick="alert(1)">Click</p>',
                      '<style>@import "https://invalid.test/style";</style>',
                      '<style>body { background: url(https://invalid.test/a) }</style>']:
            with self.subTest(extra=extra), self.assertRaises(ValueError):
                verify(self.html(extra), self.exp)
        p = self.html()
        p.write_text(p.read_text().replace('<p data-field="project_id"', '<p hidden data-field="project_id"'))
        with self.assertRaises(ValueError):
            verify(p, self.exp)
        p = self.html().rename(self.root / 'unsupported.txt')
        with self.assertRaises(ValueError):
            verify(p, self.exp)
    def test_html_financial_checks_cannot_be_self_asserted(self):
        self.fields.update(quantity='2', rate='10.005', amount='20.01')
        exp = dict(self.exp, calculations=[{'target': 'amount', 'operation': 'product',
                   'inputs': ['quantity', 'rate'], 'quantum': '0.01', 'rounding': 'ROUND_HALF_UP'}])
        self.assertEqual(verify(self.html(), exp)['checks']['calculations'], 'passed')
        self.fields['amount'] = '20.00'
        with self.assertRaisesRegex(ValueError, 'recomputed'):
            verify(self.html(), exp)
    def test_changed_during_inspection_has_no_receipt(self):
        import scripts.artifact_checks as checks
        p = self.html()
        parse = checks._html
        def concurrent_edit(raw, expectations):
            result = parse(raw, expectations)
            p.write_text(p.read_text() + '\nchanged during parsing')
            return result
        with patch.object(checks, '_html', side_effect=concurrent_edit), self.assertRaisesRegex(ValueError, 'changed'):
            verify(p, self.exp)
    def test_expectation_types_and_nonbody_sections_fail_closed(self):
        for changes in [{'schema_version': True}, {'required_text': 'Scope'},
                        {'required_text': ['']}, {'calculations': {}}, {'calculations': ''}]:
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                verify(self.html(), dict(self.exp, **changes))
        p = self.html()
        p.write_text(p.read_text().replace('<h1>Scope</h1>', '').replace('<head>', '<head><title>Scope</title>'))
        with self.assertRaises(ValueError):
            verify(p, self.exp)


if __name__ == '__main__':
    unittest.main()
