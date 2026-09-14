"""Read-only artifact checks; receipts are evidence, never approval or issuance.

Expectations v1 (unknown top-level keys fail):
* schema_version: integer 1; approved_fields: nonempty string values, including
  project_id/company/revision; required_text: list of required text fragments.
* HTML: print_css is the exact approved canonical CSS containing @media print;
  each approved field appears once as <... data-field="name">value</...> in body.
* DOCX: each approved field is an exact paragraph/table-cell line `name: value`;
  body, nested tables, all header/footer variants are read, including split runs.
* XLSX: required_sheets list, field_cells {field: 'Sheet!A1'}, cells
  {'Sheet!A1': exact JSON scalar}. Formula cells ALWAYS block, cached or not.
* Optional calculations: [{target, operation, inputs, quantum, rounding}], where
  target/inputs name semantic fields (HTML/DOCX) or Sheet!A1 addresses (XLSX).
  Operations: sum/product/subtract (subtract takes two inputs). Decimal arithmetic
  uses ROUND_HALF_UP and quantum 1/0.1/0.01/0.001/0.0001. Cover every financial
  extension/total in expectations; this checker cannot infer commercial coverage.
* Optional file_sha256 / semantic_sha256 must match readback. Semantic SHA256 is
  UTF-8 JSON of extracted approved_fields, sorted keys, separators=(',', ':'),
  ensure_ascii=False. Same fields across formats produce the same semantic hash.

No formula engine, browser rendering or PDF proof adapter is implemented. PDF
always blocks with an actionable unavailable error. Layout remains a receipt gap
for source formats; CSS checks do not prove rendering. Approval must be bound to
expectations by the caller's workflow gate; editable JSON is not authentication.
"""
import hashlib
import io
import json
import math
import re
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone
from importlib.metadata import version
from pathlib import Path
from secure_io import read_bytes, stable_io


@stable_io
def verify_artifact(path, expectations):
    """Reopen an artifact; return digest-bound receipt or raise ValueError."""
    try:
        return _verify(Path(path), expectations)
    except Exception as exc:
        raise ValueError(f'artifact unverified: {exc}') from exc


def _verify(path, expectations):
    allowed = {'schema_version', 'approved_fields', 'required_text', 'print_css',
               'file_sha256', 'semantic_sha256', 'field_cells', 'required_sheets', 'cells', 'calculations'}
    if set(expectations) - allowed or expectations.get('schema_version') != 1:
        raise ValueError('unsupported expectations schema')
    if type(expectations['schema_version']) is not int:
        raise ValueError('schema version must be integer')
    required = expectations['required_text']
    if not isinstance(required, list) or not all(isinstance(t, str) and t.strip() for t in required):
        raise ValueError('required_text must be a list of nonempty strings')
    if not isinstance(expectations.get('calculations', []), list):
        raise ValueError('calculations must be a list')
    approved = expectations['approved_fields']
    if not {'project_id', 'company', 'revision'} <= approved.keys() or not all(
            isinstance(v, str) and v.strip() for v in approved.values()):
        raise ValueError('nonempty approved identity fields required')
    path = path.absolute()
    raw = read_bytes(path)
    if path.suffix.lower() == '.pdf':
        raise ValueError('PDF text/layout proof unavailable in this verifier; retain verified source draft')
    if path.suffix.lower() == '.docx':
        fields, text, parser = _docx(raw, approved)
    elif path.suffix.lower() == '.xlsx':
        fields, text, parser = _xlsx(raw, expectations)
    elif path.suffix.lower() in ('.html', '.htm'):
        fields, text, parser = _html(raw, expectations)
    else:
        raise ValueError('unsupported artifact format')
    if fields != approved:
        raise ValueError('approved fields mismatch')
    if path.suffix.lower() != '.xlsx':
        _calculations(fields, expectations.get('calculations', []))
    if any(t not in text for t in expectations['required_text']):
        raise ValueError('required section/text missing')
    _tokens(text)
    digest = hashlib.sha256(json.dumps(fields, sort_keys=True, separators=(',', ':'),
                                     ensure_ascii=False).encode()).hexdigest()
    for key, actual in [('file_sha256', hashlib.sha256(raw).hexdigest()), ('semantic_sha256', digest)]:
        if key in expectations and expectations[key] != actual:
            raise ValueError(f'stale {key}')
    if read_bytes(path) != raw:
        raise ValueError('artifact changed during inspection')
    return {'schema_version': 1, 'project_id': fields['project_id'],
            'revision': fields['revision'], 'path': str(path),
            'file_sha256': hashlib.sha256(raw).hexdigest(), 'semantic_sha256': digest,
            'checker_version': '1.0', 'parser_version': version(parser),
            'checks': {'readback': 'passed', 'calculations': 'passed' if expectations.get('calculations') else 'not_requested'},
            'issuance': 'not_issued',
            'inspection_gaps': ['visual_layout_not_inspected'],
            'timestamp': datetime.now(timezone.utc).isoformat()}


def _tokens(text):
    if '{{' in text or '}}' in text:
        raise ValueError('unresolved tokens')


def _html(raw, expectations):
    from bs4 import BeautifulSoup
    from html.parser import HTMLParser
    class UniqueAttributes(HTMLParser):
        def handle_starttag(self, tag, attrs):
            names = [name for name, _ in attrs]
            if len(names) != len(set(names)):
                raise ValueError('duplicate HTML attribute; ambiguous browser interpretation')
    parser = UniqueAttributes()
    parser.feed(raw.decode('utf-8'))
    parser.close()
    soup = BeautifulSoup(raw.decode('utf-8'), 'html.parser')
    if soup.find(['script', 'link', 'iframe', 'object', 'embed']) or soup.select('[src]'):
        raise ValueError('scripts or resource dependencies are forbidden')
    _tokens(str(soup))
    css = expectations['print_css']
    if '@media print' not in css or css not in '\n'.join(n.get_text() for n in soup.find_all('style')):
        raise ValueError('required canonical print CSS missing')
    styles = '\n'.join(n.get_text() for n in soup.find_all('style'))
    styles += '\n'.join(str(n.get('style', '')) for n in soup.find_all(True))
    if re.search(r'@import|url\s*\(|expression\s*\(', styles, re.I):
        raise ValueError('active CSS/resource dependency')
    visible_css = re.sub(r'/\*.*?\*/', '', styles, flags=re.S)
    if re.search(r'display\s*:\s*none|(?:content-)?visibility\s*:\s*(?:hidden|collapse)'
                 r'|opacity\s*:\s*0(?:\.0*)?\s*(?:[;!}]|$)', visible_css, re.I):
        raise ValueError('hidden content CSS is forbidden, including canonical CSS')
    url_attrs = {'href', 'xlink:href', 'src', 'action', 'formaction', 'poster',
                 'background', 'cite', 'data', 'longdesc', 'usemap', 'profile', 'manifest'}
    for node in soup.find_all(True):
        if node.name == 'meta' and str(node.get('http-equiv', '')).strip().lower() == 'refresh':
            raise ValueError('meta refresh is forbidden')
        for key in url_attrs & node.attrs.keys():
            # HTML entities are decoded by the parser; strip URL control whitespace.
            url = re.sub(r'[\x00-\x20]', '', str(node[key])).lower()
            if url.startswith(('javascript:', 'vbscript:', 'data:')):
                raise ValueError('active URL scheme is forbidden')
        if (node.name in {'form', 'input', 'button', 'select', 'textarea'}
                or {'action', 'formaction'} & node.attrs.keys()):
            raise ValueError('HTML submission controls are forbidden')
        if node.name == 'template' or (node.name == 'dialog' and 'open' not in node.attrs):
            raise ValueError('hidden template or closed dialog content')
        if node.name == 'details' and 'open' not in node.attrs:
            summary = node.find('summary', recursive=False)
            if any(child is not summary and str(child).strip() for child in node.contents):
                raise ValueError('hidden closed details content outside first summary')
        if ('hidden' in node.attrs or 'inert' in node.attrs or 'popover' in node.attrs
                or str(node.get('aria-hidden', '')).strip().lower() == 'true'
                or any(key.lower().startswith('on') for key in node.attrs)):
            raise ValueError('hidden content or active event handler')
    body = soup.body
    if body is None:
        raise ValueError('HTML body missing')
    fields = {}
    for node in body.select('[data-field]'):
        name = node['data-field']
        if name in fields:
            raise ValueError('duplicate semantic field')
        fields[name] = node.get_text(' ', strip=True)
    return fields, body.get_text(' ', strip=True), 'beautifulsoup4'


def _docx(raw, approved):
    from docx import Document
    doc = Document(io.BytesIO(raw))
    roots = [doc.element]
    for section in doc.sections:
        for name in ('header', 'footer', 'first_page_header', 'first_page_footer',
                     'even_page_header', 'even_page_footer'):
            roots.append(getattr(section, name)._element)
    lines = []
    for root in roots:
        for para in root.xpath('.//w:p'):
            lines.append(''.join(para.xpath('.//w:t/text()')))
    fields = {}
    for name in approved:
        matches = [line.split(':', 1)[1].strip() for line in lines if line.startswith(name + ':')]
        if not matches or any(value != approved[name] for value in matches):
            raise ValueError(f'approved DOCX field mismatch: {name}')
        fields[name] = matches[0]
    return fields, '\n'.join(lines), 'python-docx'


def _xlsx(raw, expectations):
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(raw), data_only=False)
    try:
        if not set(expectations['required_sheets']) <= set(book.sheetnames):
            raise ValueError('required sheet missing')
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if cell.data_type == 'f':
                        raise ValueError(f'unsupported formula: {sheet.title}!{cell.coordinate}; export reviewed values')
        values = {f'{sheet.title}!{cell.coordinate}': cell.value
                  for sheet in book for row in sheet for cell in row if cell.value is not None}
        for address, expected in expectations['cells'].items():
            actual = values.get(address)
            # JSON booleans are distinct from numbers, even though True == 1 in Python.
            if isinstance(actual, bool) != isinstance(expected, bool) or actual != expected:
                raise ValueError(f'cell mismatch: {address}')
        fields = {name: str(values[address]) for name, address in expectations['field_cells'].items()}
        _calculations(values, expectations.get('calculations', []))
        return fields, '\n'.join(map(str, values.values())), 'openpyxl'
    finally:
        book.close()


def _number(value):
    if isinstance(value, bool):
        raise ValueError('boolean is not a financial number')
    result = Decimal(str(value))
    if not result.is_finite():
        raise ValueError('nonfinite financial number')
    return result


def _calculations(values, calculations):
    for calc in calculations:
        if set(calc) != {'target', 'operation', 'inputs', 'quantum', 'rounding'}:
            raise ValueError('invalid calculation fields')
        if calc['rounding'] != 'ROUND_HALF_UP' or calc['quantum'] not in ('1', '0.1', '0.01', '0.001', '0.0001'):
            raise ValueError('unsupported explicit rounding basis')
        inputs = [_number(values[address]) for address in calc['inputs']]
        if not inputs:
            raise ValueError('calculation inputs required')
        operation = calc['operation']
        if operation == 'sum':
            result = sum(inputs, Decimal(0))
        elif operation == 'product':
            result = math.prod(inputs)
        elif operation == 'subtract' and len(inputs) == 2:
            result = inputs[0] - inputs[1]
        else:
            raise ValueError('unsupported calculation')
        result = result.quantize(Decimal(calc['quantum']), rounding=ROUND_HALF_UP)
        if _number(values[calc['target']]) != result:
            raise ValueError(f'recomputed total mismatch: {calc["target"]}')
